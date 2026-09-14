"""
Phase 3 reconnaissance: check whether the Service Charter (.doc) holds
genuine structured content (a service-standards table) worth
extracting, or is narrative-only like the Acts.
"""
from pathlib import Path

PATH = Path("data/raw/other/SRC-GOVE-009_LIVINGSTONE-CITY-COUNCIL-SERVICE-CHARTER.doc")

try:
    import docx
except ImportError:
    print("python-docx is NOT installed.")
    raise SystemExit(1)

try:
    d = docx.Document(str(PATH))
except Exception as e:
    print(f"FAILED to open as .docx: {type(e).__name__}: {e}")
    print("Likely a genuine legacy .doc (binary format) - python-docx cannot read it.")
    raise SystemExit(0)

print(f"Paragraphs: {len(d.paragraphs)}")
print(f"Tables detected: {len(d.tables)}")

print("\nFirst 10 non-empty paragraphs:")
shown = 0
for p in d.paragraphs:
    if p.text.strip():
        print(f"  {p.text.strip()[:150]}")
        shown += 1
    if shown >= 10:
        break

if d.tables:
    print(f"\nFirst table preview (table 1 of {len(d.tables)}):")
    t = d.tables[0]
    for row in t.rows[:5]:
        print("  | " + " | ".join(cell.text.strip()[:30] for cell in row.cells))
